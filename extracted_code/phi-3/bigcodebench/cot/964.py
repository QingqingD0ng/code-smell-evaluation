import os
from pathlib import Path
import pandas as pd
from docx import Document
from openpyxl import load_workbook

def task_func(source_directory: str, target_directory: str) -> int:
    source_path = Path(source_directory)
    target_path = Path(target_directory)
    target_path.mkdir(parents=True, exist_ok=True)
    converted_count = 0

    for file_path in source_path.rglob('*'):
        if file_path.suffix.lower() in ['.txt', '.docx', '.xlsx', '.csv']:
            target_file_path = target_path / f"{file_path.stem}.csv"
            if file_path.suffix.lower() == '.txt':
                with open(file_path, 'r', encoding='utf-8') as file:
                    text = file.read()
                    pd.DataFrame([text]).to_csv(target_file_path, index=False)
            elif file_path.suffix.lower() == '.docx':
                doc = Document(file_path)
                text = '\n'.join([para.text for para in doc.paragraphs])
                pd.DataFrame([text]).to_csv(target_file_path, index=False)
            elif file_path.suffix.lower() == '.xlsx':
                workbook = load_workbook(file_path)
                sheet = workbook.active
                text = '\n'.join([cell.value for cell in sheet[1]])
                pd.DataFrame([text]).to_csv(target_file_path, index=False)
            elif file_path.suffix.lower() == '.csv':
                continue
            else:
                raise ValueError("Unsupported file type")
            os.replace(target_file_path, file_path)
            converted_count += 1

    return converted_count

# Example usage:
# converted_files_count = task_func('/Users/test/Documents', '/Users/test/Documents/csv_files')