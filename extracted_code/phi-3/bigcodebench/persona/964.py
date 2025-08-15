import os

from pathlib import Path

import pandas as pd

from docx import Document

from openpyxl import load_workbook


def task_func(source_directory: str, target_directory: str) -> int:
    source_path = Path(source_directory)
    target_path = Path(target_directory)
    target_path.mkdir(parents=True, exist_ok=True)
    converted_count = 0

    for file in source_path.glob('*'):
        if file.suffix in ['.txt', '.docx', '.xlsx', '.csv']:
            target_file = target_path / (file.stem + '.csv')
            if file.suffix == '.txt':
                with open(file, 'r') as txt_file:
                    text = txt_file.read()
            elif file.suffix == '.docx':
                doc = Document(file)
                text = '\n'.join([para.text for para in doc.paragraphs])
            elif file.suffix == '.xlsx':
                wb = load_workbook(file)
                text = '\n'.join([cell.value for sheet in wb.worksheets for cell in sheet])
            elif file.suffix == '.csv':
                df = pd.read_csv(file)
                text = df['Text'].str.cat(sep='\n')
            pd.DataFrame({'Text': [text]}).to_csv(target_file, index=False)
            converted_count += 1

    return converted_count